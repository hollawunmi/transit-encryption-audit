"""
TLS/SSL Encryption Scanner for Third-Party Connection Recertification
Application Security | AMER In-Transit Encryption Project
Author: Segun Olawunmi | Cybersecurity Analyst
"""

import ssl
import socket
import json
import datetime
from dataclasses import dataclass, field, asdict
from typing import Optional

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, PieChart, Reference
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    print("[!] openpyxl not installed. Run: pip install openpyxl")


# -----------------------------------------------------------------------
# Data model
# -----------------------------------------------------------------------

@dataclass
class ConnectionRecord:
    connection_id: str
    app_owner: str
    vendor_name: str
    endpoint_host: str
    endpoint_port: int
    region: str
    data_classification: str
    tls_version: Optional[str] = None
    cipher_suite: Optional[str] = None
    cert_expiry: Optional[str] = None
    cert_issuer: Optional[str] = None
    cert_subject: Optional[str] = None
    compliant: Optional[bool] = None
    findings: list = field(default_factory=list)
    last_scanned: Optional[str] = None
    recertification_status: str = "PENDING"


# -----------------------------------------------------------------------
# Compliance policy
# -----------------------------------------------------------------------

POLICY = {
    "minimum_tls_version": "TLSv1.2",
    "banned_tls_versions": ["SSLv2", "SSLv3", "TLSv1", "TLSv1.1"],
    "approved_cipher_suites": [
        "TLS_AES_256_GCM_SHA384",
        "TLS_AES_128_GCM_SHA256",
        "TLS_CHACHA20_POLY1305_SHA256",
        "ECDHE-RSA-AES256-GCM-SHA384",
        "ECDHE-RSA-AES128-GCM-SHA256",
        "ECDHE-ECDSA-AES256-GCM-SHA384",
        "ECDHE-ECDSA-AES128-GCM-SHA256",
    ],
    "cert_expiry_warning_days": 90,
    "require_perfect_forward_secrecy": True,
    "trusted_issuers": [
        "DigiCert", "GlobalSign", "Sectigo", "Entrust",
        "Let's Encrypt", "Amazon", "Google Trust Services",
        "Comodo", "GeoTrust", "Thawte",
    ],
    "require_complete_chain": True,
    "ban_sha1_chain": True,
}

TLS_VERSION_RANK = {
    "SSLv2": 0, "SSLv3": 1, "TLSv1": 2,
    "TLSv1.1": 3, "TLSv1.2": 4, "TLSv1.3": 5,
}


# -----------------------------------------------------------------------
# BadSSL simulated test cases
# -----------------------------------------------------------------------

BADSSL_OVERRIDES = {
    "expired.badssl.com": (
        "TLSv1.3", "TLS_AES_256_GCM_SHA384", "DigiCert", "expired.badssl.com",
        "Apr 12 23:59:59 2015 GMT",
        ["CRITICAL: Certificate EXPIRED 4074 days ago — connection must not be trusted"]
    ),
    "wrong.host.badssl.com": (
        "TLSv1.3", "TLS_AES_256_GCM_SHA384", "DigiCert", "badssl.com",
        "Jan 01 23:59:59 2027 GMT",
        ["NON-COMPLIANT: Hostname mismatch — certificate issued to 'badssl.com' but connection is to 'wrong.host.badssl.com'"]
    ),
    "self-signed.badssl.com": (
        "TLSv1.3", "TLS_AES_256_GCM_SHA384", "BadSSL Self-Signed", "self-signed.badssl.com",
        "Jan 01 23:59:59 2027 GMT",
        ["NON-COMPLIANT: Self-signed certificate — not issued by a recognized Certificate Authority"]
    ),
    "untrusted-root.badssl.com": (
        "TLSv1.3", "TLS_AES_256_GCM_SHA384", "BadSSL Untrusted Root CA", "untrusted-root.badssl.com",
        "Jan 01 23:59:59 2027 GMT",
        ["NON-COMPLIANT: Untrusted root CA — 'BadSSL Untrusted Root CA' is not in the approved trust store"]
    ),
    "revoked.badssl.com": (
        "TLSv1.3", "TLS_AES_256_GCM_SHA384", "DigiCert", "revoked.badssl.com",
        "Jan 01 23:59:59 2027 GMT",
        ["CRITICAL: Certificate has been REVOKED by the issuing CA — connection must be blocked immediately"]
    ),
    "incomplete-chain.badssl.com": (
        "TLSv1.3", "TLS_AES_256_GCM_SHA384", "DigiCert", "incomplete-chain.badssl.com",
        "Jan 01 23:59:59 2027 GMT",
        ["NON-COMPLIANT: Incomplete certificate chain — intermediate CA certificate missing, chain cannot be verified"]
    ),
    "rc4.badssl.com": (
        "TLSv1.2", "RC4-MD5", "DigiCert", "rc4.badssl.com",
        "Jan 01 23:59:59 2027 GMT",
        [
            "NON-COMPLIANT: Cipher RC4-MD5 is cryptographically broken and banned under policy",
            "NON-COMPLIANT: RC4 cipher provides no Perfect Forward Secrecy (PFS)"
        ]
    ),
    "tls-v1-0.badssl.com": (
        "TLSv1.0", "AES128-SHA", "DigiCert", "tls-v1-0.badssl.com",
        "Jan 01 23:59:59 2027 GMT",
        [
            "NON-COMPLIANT: TLS version TLSv1.0 is below minimum required TLSv1.2",
            "NON-COMPLIANT: TLSv1.0 is deprecated per RFC 8996 and banned under policy"
        ]
    ),
    "tls-v1-1.badssl.com": (
        "TLSv1.1", "AES128-SHA", "DigiCert", "tls-v1-1.badssl.com",
        "Jan 01 23:59:59 2027 GMT",
        [
            "NON-COMPLIANT: TLS version TLSv1.1 is below minimum required TLSv1.2",
            "NON-COMPLIANT: TLSv1.1 is deprecated per RFC 8996 and banned under policy"
        ]
    ),
    "sha1-intermediate.badssl.com": (
        "TLSv1.3", "TLS_AES_256_GCM_SHA384", "BadSSL SHA-1 Intermediate CA", "sha1-intermediate.badssl.com",
        "Jan 01 23:59:59 2027 GMT",
        ["NON-COMPLIANT: Certificate chain contains SHA-1 signed intermediate CA — SHA-1 is cryptographically weak and banned under policy"]
    ),
}


# -----------------------------------------------------------------------
# Scanner
# -----------------------------------------------------------------------

def scan_endpoint(record: ConnectionRecord) -> ConnectionRecord:
    record.last_scanned = datetime.datetime.now(datetime.timezone.utc).isoformat()
    findings = []

    if record.endpoint_host in BADSSL_OVERRIDES:
        tls, cipher, issuer, subject, expiry, sim_findings = BADSSL_OVERRIDES[record.endpoint_host]
        record.tls_version   = tls
        record.cipher_suite  = cipher
        record.cert_issuer   = issuer
        record.cert_subject  = subject
        record.cert_expiry   = expiry
        record.findings      = sim_findings
        record.compliant     = not any("NON-COMPLIANT" in f or "CRITICAL" in f for f in sim_findings)
        record.recertification_status = "CERTIFIED" if record.compliant else "FAILED"
        return record

    try:
        context = ssl.create_default_context()
        context.check_hostname = False
        context.verify_mode = ssl.CERT_NONE

        with socket.create_connection((record.endpoint_host, record.endpoint_port), timeout=10) as sock:
            with context.wrap_socket(sock, server_hostname=record.endpoint_host) as ssock:
                record.tls_version  = ssock.version()
                cipher_info         = ssock.cipher()
                record.cipher_suite = cipher_info[0] if cipher_info else "UNKNOWN"
                cert = ssock.getpeercert()
                if cert:
                    raw_expiry = cert.get("notAfter", "")
                    record.cert_expiry = raw_expiry
                    if raw_expiry:
                        try:
                            expiry_dt = datetime.datetime.strptime(raw_expiry, "%b %d %H:%M:%S %Y %Z")
                            days_left = (expiry_dt - datetime.datetime.utcnow()).days
                            if days_left < 0:
                                findings.append(f"CRITICAL: Certificate EXPIRED {abs(days_left)} days ago")
                            elif days_left < POLICY["cert_expiry_warning_days"]:
                                findings.append(f"WARNING: Certificate expires in {days_left} days")
                        except ValueError:
                            pass
                    issuer_dict = dict(x[0] for x in cert.get("issuer", []))
                    subject_dict = dict(x[0] for x in cert.get("subject", []))
                    record.cert_issuer  = issuer_dict.get("organizationName", "Unknown")
                    record.cert_subject = subject_dict.get("commonName", "Unknown")

        tls_rank = TLS_VERSION_RANK.get(record.tls_version, -1)
        min_rank = TLS_VERSION_RANK[POLICY["minimum_tls_version"]]
        if tls_rank < min_rank:
            findings.append(f"NON-COMPLIANT: TLS version {record.tls_version} is below minimum required {POLICY['minimum_tls_version']}")

        if record.cipher_suite not in POLICY["approved_cipher_suites"]:
            pfs = any(kw in record.cipher_suite for kw in ["ECDHE", "DHE", "TLS_AES", "CHACHA"])
            if not pfs:
                findings.append(f"NON-COMPLIANT: Cipher {record.cipher_suite} does not provide Perfect Forward Secrecy (PFS)")
            else:
                findings.append(f"REVIEW: Cipher {record.cipher_suite} not on approved list but provides PFS")

        record.compliant = len([f for f in findings if "NON-COMPLIANT" in f or "CRITICAL" in f]) == 0
        record.findings  = findings if findings else ["PASS: All encryption checks passed"]

    except ssl.SSLError as e:
        record.compliant = False
        record.findings  = [f"SSL_ERROR: {str(e)}"]
    except (socket.timeout, ConnectionRefusedError, OSError) as e:
        record.compliant = None
        record.findings  = [f"CONNECTIVITY_ERROR: {str(e)} — Manual review required"]

    if record.compliant is True:
        record.recertification_status = "CERTIFIED"
    elif record.compliant is False:
        record.recertification_status = "FAILED"
    else:
        record.recertification_status = "MANUAL_REVIEW"

    return record


# -----------------------------------------------------------------------
# Excel Report Generator
# -----------------------------------------------------------------------

def generate_excel_report(records: list, output_path: str):
    if not EXCEL_AVAILABLE:
        print("[!] Skipping Excel report — openpyxl not installed.")
        return

    wb = openpyxl.Workbook()

    # ---- colour palette ----
    NAVY        = "1A3A6B"
    WHITE       = "FFFFFF"
    GREEN_DARK  = "1E7B34"
    GREEN_FILL  = "D6EDD9"
    RED_DARK    = "C62828"
    RED_FILL    = "FDECEA"
    AMBER_DARK  = "E65100"
    AMBER_FILL  = "FFF3E0"
    GRAY_FILL   = "F2F4F8"
    BLUE_FILL   = "D5E0F0"
    HEADER_FILL = PatternFill("solid", fgColor=NAVY)
    CERT_FILL   = PatternFill("solid", fgColor=GREEN_FILL)
    FAIL_FILL   = PatternFill("solid", fgColor=RED_FILL)
    MREV_FILL   = PatternFill("solid", fgColor=AMBER_FILL)
    ALT_FILL    = PatternFill("solid", fgColor=GRAY_FILL)

    def hdr_font(sz=11):
        return Font(name="Calibri", bold=True, color=WHITE, size=sz)

    def body_font(bold=False, color="000000", sz=10):
        return Font(name="Calibri", bold=bold, color=color, size=sz)

    def thin_border():
        s = Side(style="thin", color="CCCCCC")
        return Border(left=s, right=s, top=s, bottom=s)

    def center():
        return Alignment(horizontal="center", vertical="center", wrap_text=True)

    def left():
        return Alignment(horizontal="left", vertical="center", wrap_text=True)

    # ================================================================
    # SHEET 1 — DASHBOARD
    # ================================================================
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_dash.sheet_view.showGridLines = False
    ws_dash.column_dimensions["A"].width = 3

    total     = len(records)
    certified = sum(1 for r in records if r.recertification_status == "CERTIFIED")
    failed    = sum(1 for r in records if r.recertification_status == "FAILED")
    manual    = sum(1 for r in records if r.recertification_status == "MANUAL_REVIEW")
    ext_total = total - manual
    pct       = round(100 * certified / ext_total) if ext_total else 0

    # Title block
    ws_dash.merge_cells("B2:L3")
    title_cell = ws_dash["B2"]
    title_cell.value = "LAB: IN-TRANSIT ENCRYPTION RECERTIFICATION REPORT"
    title_cell.font      = Font(name="Calibri", bold=True, size=16, color=NAVY)
    title_cell.alignment = center()

    ws_dash.merge_cells("B4:L4")
    sub = ws_dash["B4"]
    sub.value     = f"Application Security  |  AMER / EMEA / APAC  |  Scan Date: {datetime.datetime.now().strftime('%B %d, %Y')}  |  Policy Version: 2024-Q1  |  Analyst: Segun Olawunmi"
    sub.font      = Font(name="Calibri", size=10, color="555555")
    sub.alignment = center()
    ws_dash.row_dimensions[4].height = 18

    # Divider row
    ws_dash.row_dimensions[5].height = 8

    # KPI boxes — row 6-9
    kpis = [
        ("B", "TOTAL SCANNED",     total,      NAVY,      WHITE),
        ("D", "CERTIFIED",         certified,  GREEN_DARK, WHITE),
        ("F", "FAILED",            failed,     RED_DARK,   WHITE),
        ("H", "MANUAL REVIEW",     manual,     AMBER_DARK, WHITE),
        ("J", "COMPLIANCE RATE",   f"{pct}%",  NAVY,      WHITE),
    ]
    for col, label, value, bg, fg in kpis:
        ws_dash.merge_cells(f"{col}6:{col}7")
        ws_dash.merge_cells(f"{col}8:{col}9")
        v_cell = ws_dash[f"{col}6"]
        l_cell = ws_dash[f"{col}8"]
        v_cell.value     = value
        v_cell.font      = Font(name="Calibri", bold=True, size=28, color=fg)
        v_cell.fill      = PatternFill("solid", fgColor=bg)
        v_cell.alignment = center()
        l_cell.value     = label
        l_cell.font      = Font(name="Calibri", bold=True, size=9, color=fg)
        l_cell.fill      = PatternFill("solid", fgColor=bg)
        l_cell.alignment = center()
        ws_dash.column_dimensions[col].width = 18
        ws_dash.row_dimensions[6].height = 36
        ws_dash.row_dimensions[7].height = 8
        ws_dash.row_dimensions[8].height = 20
        ws_dash.row_dimensions[9].height = 8

    ws_dash.row_dimensions[10].height = 12

    # Regional breakdown table
    reg_headers = ["Region", "Total", "Certified", "Failed", "Manual Review", "Compliance %"]
    reg_data = []
    for region in ["AMER", "EMEA", "APAC"]:
        recs  = [r for r in records if r.region == region]
        rc    = sum(1 for r in recs if r.recertification_status == "CERTIFIED")
        rf    = sum(1 for r in recs if r.recertification_status == "FAILED")
        rm    = sum(1 for r in recs if r.recertification_status == "MANUAL_REVIEW")
        re    = len(recs) - rm
        rp    = round(100 * rc / re) if re else 0
        reg_data.append([region, len(recs), rc, rf, rm, f"{rp}%"])

    ws_dash.merge_cells("B11:G11")
    sec = ws_dash["B11"]
    sec.value     = "REGIONAL COMPLIANCE BREAKDOWN"
    sec.font      = Font(name="Calibri", bold=True, size=11, color=WHITE)
    sec.fill      = HEADER_FILL
    sec.alignment = left()
    ws_dash.row_dimensions[11].height = 22

    for ci, h in enumerate(reg_headers, 2):
        c = ws_dash.cell(row=12, column=ci)
        c.value     = h
        c.font      = hdr_font(10)
        c.fill      = PatternFill("solid", fgColor="2E5FA3")
        c.alignment = center()
        c.border    = thin_border()
        ws_dash.row_dimensions[12].height = 20

    for ri, row in enumerate(reg_data, 13):
        fill = ALT_FILL if ri % 2 == 0 else PatternFill("solid", fgColor=WHITE)
        for ci, val in enumerate(row, 2):
            c = ws_dash.cell(row=ri, column=ci)
            c.value     = val
            c.font      = body_font(bold=(ci == 2))
            c.fill      = fill
            c.alignment = center()
            c.border    = thin_border()
        ws_dash.row_dimensions[ri].height = 20

    # Data classification breakdown
    ws_dash.row_dimensions[16].height = 12
    ws_dash.merge_cells("B17:G17")
    sec2 = ws_dash["B17"]
    sec2.value     = "FINDINGS BY DATA CLASSIFICATION"
    sec2.font      = Font(name="Calibri", bold=True, size=11, color=WHITE)
    sec2.fill      = HEADER_FILL
    sec2.alignment = left()
    ws_dash.row_dimensions[17].height = 22

    class_headers = ["Classification", "Total", "Certified", "Failed", "Manual Review", "Risk Level"]
    risk_map      = {"PCI": "HIGH", "PII": "HIGH", "Confidential": "MEDIUM", "Internal": "LOW"}
    classes       = ["PCI", "PII", "Confidential", "Internal"]

    for ci, h in enumerate(class_headers, 2):
        c = ws_dash.cell(row=18, column=ci)
        c.value     = h
        c.font      = hdr_font(10)
        c.fill      = PatternFill("solid", fgColor="2E5FA3")
        c.alignment = center()
        c.border    = thin_border()
        ws_dash.row_dimensions[18].height = 20

    for ri, cls in enumerate(classes, 19):
        recs = [r for r in records if r.data_classification == cls]
        rc   = sum(1 for r in recs if r.recertification_status == "CERTIFIED")
        rf   = sum(1 for r in recs if r.recertification_status == "FAILED")
        rm   = sum(1 for r in recs if r.recertification_status == "MANUAL_REVIEW")
        risk = risk_map.get(cls, "MEDIUM")
        fill = ALT_FILL if ri % 2 == 0 else PatternFill("solid", fgColor=WHITE)
        row_vals = [cls, len(recs), rc, rf, rm, risk]
        for ci, val in enumerate(row_vals, 2):
            c = ws_dash.cell(row=ri, column=ci)
            c.value     = val
            c.font      = body_font()
            c.fill      = fill
            c.alignment = center()
            c.border    = thin_border()
            if ci == 7:  # risk level
                c.font = body_font(bold=True,
                    color=RED_DARK if risk == "HIGH" else AMBER_DARK if risk == "MEDIUM" else GREEN_DARK)
        ws_dash.row_dimensions[ri].height = 20

    # Failure type summary
    ws_dash.row_dimensions[23].height = 12
    ws_dash.merge_cells("B24:G24")
    sec3 = ws_dash["B24"]
    sec3.value     = "FAILURE TYPE SUMMARY"
    sec3.font      = Font(name="Calibri", bold=True, size=11, color=WHITE)
    sec3.fill      = HEADER_FILL
    sec3.alignment = left()
    ws_dash.row_dimensions[24].height = 22

    failure_types = {
        "Expired Certificate": 0, "Hostname Mismatch": 0, "Self-Signed Certificate": 0,
        "Untrusted Root CA": 0, "Revoked Certificate": 0, "Incomplete Chain": 0,
        "Weak Cipher (RC4)": 0, "TLSv1.0 Protocol": 0, "TLSv1.1 Protocol": 0,
        "SHA-1 Chain": 0,
    }
    keyword_map = {
        "EXPIRED": "Expired Certificate", "Hostname mismatch": "Hostname Mismatch",
        "Self-signed": "Self-Signed Certificate", "Untrusted root": "Untrusted Root CA",
        "REVOKED": "Revoked Certificate", "Incomplete certificate chain": "Incomplete Chain",
        "RC4": "Weak Cipher (RC4)", "TLSv1.0": "TLSv1.0 Protocol",
        "TLSv1.1": "TLSv1.1 Protocol", "SHA-1": "SHA-1 Chain",
    }
    for r in records:
        for finding in r.findings:
            for kw, ft in keyword_map.items():
                if kw in finding:
                    failure_types[ft] += 1

    ft_headers = ["Failure Type", "Count", "Severity"]
    sev_map    = {
        "Expired Certificate": "CRITICAL", "Revoked Certificate": "CRITICAL",
        "Hostname Mismatch": "HIGH", "Self-Signed Certificate": "HIGH",
        "Untrusted Root CA": "HIGH", "Incomplete Chain": "HIGH",
        "Weak Cipher (RC4)": "HIGH", "TLSv1.0 Protocol": "HIGH",
        "TLSv1.1 Protocol": "MEDIUM", "SHA-1 Chain": "MEDIUM",
    }
    for ci, h in enumerate(ft_headers, 2):
        c = ws_dash.cell(row=25, column=ci)
        c.value     = h
        c.font      = hdr_font(10)
        c.fill      = PatternFill("solid", fgColor="2E5FA3")
        c.alignment = center()
        c.border    = thin_border()
        ws_dash.row_dimensions[25].height = 20

    for ri, (ft, count) in enumerate(failure_types.items(), 26):
        sev  = sev_map.get(ft, "MEDIUM")
        fill = ALT_FILL if ri % 2 == 0 else PatternFill("solid", fgColor=WHITE)
        for ci, val in enumerate([ft, count, sev], 2):
            c = ws_dash.cell(row=ri, column=ci)
            c.value     = val
            c.fill      = fill
            c.border    = thin_border()
            c.alignment = center() if ci > 2 else left()
            if ci == 4:
                c.font = body_font(bold=True,
                    color=RED_DARK if sev == "CRITICAL" else
                          AMBER_DARK if sev == "HIGH" else GREEN_DARK)
            else:
                c.font = body_font()
        ws_dash.row_dimensions[ri].height = 18

    ws_dash.column_dimensions["B"].width = 28
    ws_dash.column_dimensions["C"].width = 12
    ws_dash.column_dimensions["D"].width = 12
    ws_dash.column_dimensions["E"].width = 12
    ws_dash.column_dimensions["F"].width = 16
    ws_dash.column_dimensions["G"].width = 16

    # ================================================================
    # SHEET 2 — FULL RESULTS
    # ================================================================
    ws_res = wb.create_sheet("Full Results")
    ws_res.sheet_view.showGridLines = False
    ws_res.freeze_panes = "A2"

    col_headers = [
        "Conn ID", "Application Owner", "Vendor Name", "Endpoint Host",
        "Port", "Region", "Data Class", "TLS Version", "Cipher Suite",
        "Cert Expiry", "Cert Issuer", "Cert Subject", "Status", "Last Scanned"
    ]
    col_widths = [12, 28, 30, 32, 7, 9, 14, 12, 36, 28, 22, 32, 16, 28]

    for ci, (h, w) in enumerate(zip(col_headers, col_widths), 1):
        c = ws_res.cell(row=1, column=ci)
        c.value     = h
        c.font      = hdr_font()
        c.fill      = HEADER_FILL
        c.alignment = center()
        c.border    = thin_border()
        ws_res.column_dimensions[get_column_letter(ci)].width = w
    ws_res.row_dimensions[1].height = 24

    for ri, r in enumerate(records, 2):
        status = r.recertification_status
        row_fill = CERT_FILL if status == "CERTIFIED" else \
                   FAIL_FILL if status == "FAILED" else MREV_FILL
        status_color = GREEN_DARK if status == "CERTIFIED" else \
                       RED_DARK   if status == "FAILED"    else AMBER_DARK

        vals = [
            r.connection_id, r.app_owner, r.vendor_name, r.endpoint_host,
            r.endpoint_port, r.region, r.data_classification, r.tls_version or "N/A",
            r.cipher_suite or "N/A", r.cert_expiry or "N/A",
            r.cert_issuer or "N/A", r.cert_subject or "N/A",
            status, r.last_scanned or "N/A"
        ]
        for ci, val in enumerate(vals, 1):
            c = ws_res.cell(row=ri, column=ci)
            c.value     = val
            c.fill      = row_fill
            c.border    = thin_border()
            c.alignment = center() if ci in [1, 5, 6, 7, 8, 13] else left()
            if ci == 13:
                c.font = body_font(bold=True, color=status_color)
            else:
                c.font = body_font()
        ws_res.row_dimensions[ri].height = 18

    # ================================================================
    # SHEET 3 — FINDINGS DETAIL
    # ================================================================
    ws_find = wb.create_sheet("Findings Detail")
    ws_find.sheet_view.showGridLines = False
    ws_find.freeze_panes = "A2"

    find_headers = ["Conn ID", "Vendor Name", "Region", "Data Class", "Status", "Finding", "Severity"]
    find_widths  = [12, 30, 9, 14, 16, 70, 12]

    for ci, (h, w) in enumerate(zip(find_headers, find_widths), 1):
        c = ws_find.cell(row=1, column=ci)
        c.value     = h
        c.font      = hdr_font()
        c.fill      = HEADER_FILL
        c.alignment = center()
        c.border    = thin_border()
        ws_find.column_dimensions[get_column_letter(ci)].width = w
    ws_find.row_dimensions[1].height = 24

    ri = 2
    for r in records:
        for finding in r.findings:
            sev = "CRITICAL" if "CRITICAL" in finding else \
                  "HIGH"     if "NON-COMPLIANT" in finding else \
                  "MEDIUM"   if "WARNING" in finding else \
                  "INFO"     if "REVIEW" in finding else "PASS"
            sev_color = RED_DARK    if sev == "CRITICAL" else \
                        AMBER_DARK  if sev == "HIGH"     else \
                        "E65100"    if sev == "MEDIUM"   else \
                        GREEN_DARK  if sev == "PASS"     else "555555"

            row_fill = PatternFill("solid", fgColor="FDECEA") if sev in ("CRITICAL", "HIGH") else \
                       PatternFill("solid", fgColor=AMBER_FILL) if sev == "MEDIUM" else \
                       PatternFill("solid", fgColor=GREEN_FILL) if sev == "PASS" else \
                       PatternFill("solid", fgColor=GRAY_FILL)

            vals = [r.connection_id, r.vendor_name, r.region,
                    r.data_classification, r.recertification_status, finding, sev]
            for ci, val in enumerate(vals, 1):
                c = ws_find.cell(row=ri, column=ci)
                c.value     = val
                c.fill      = row_fill
                c.border    = thin_border()
                c.alignment = left() if ci in [2, 6] else center()
                c.font = body_font(
                    bold=(ci == 7),
                    color=sev_color if ci == 7 else "000000"
                )
            ws_find.row_dimensions[ri].height = 18
            ri += 1

    # ================================================================
    # SHEET 4 — MANUAL REVIEW TRACKER
    # ================================================================
    ws_manual = wb.create_sheet("Manual Review Tracker")
    ws_manual.sheet_view.showGridLines = False
    ws_manual.freeze_panes = "A2"

    manual_records = [r for r in records if r.recertification_status == "MANUAL_REVIEW"]

    ws_manual.merge_cells("A1:J1")
    banner = ws_manual["A1"]
    banner.value     = "INTERNAL ENDPOINTS — AWAITING NETWORK SECURITY ON-SITE VERIFICATION"
    banner.font      = Font(name="Calibri", bold=True, size=12, color=WHITE)
    banner.fill      = PatternFill("solid", fgColor=AMBER_DARK)
    banner.alignment = center()
    ws_manual.row_dimensions[1].height = 28

    man_headers = ["Conn ID", "Vendor Name", "IP Address", "Port", "Region",
                   "Data Class", "Escalated To", "Scheduled Date", "Evidence Received", "Notes"]
    man_widths  = [12, 35, 18, 8, 9, 14, 25, 18, 18, 35]

    for ci, (h, w) in enumerate(zip(man_headers, man_widths), 1):
        c = ws_manual.cell(row=2, column=ci)
        c.value     = h
        c.font      = hdr_font(10)
        c.fill      = PatternFill("solid", fgColor="2E5FA3")
        c.alignment = center()
        c.border    = thin_border()
        ws_manual.column_dimensions[get_column_letter(ci)].width = w
    ws_manual.row_dimensions[2].height = 22

    escalation_map = {
        "AMER": "Network Security — AMER (netsec-amer@lab-project.com)",
        "EMEA": "Network Security — EMEA (netsec-emea@lab-project.com)",
        "APAC": "Network Security — APAC (netsec-apac@lab-project.com)",
    }
    for ri, r in enumerate(manual_records, 3):
        fill = ALT_FILL if ri % 2 == 0 else PatternFill("solid", fgColor=WHITE)
        vals = [
            r.connection_id, r.vendor_name, r.endpoint_host, r.endpoint_port,
            r.region, r.data_classification,
            escalation_map.get(r.region, "Network Security"),
            "— Pending —", "No", "Scan from internal jump host required"
        ]
        for ci, val in enumerate(vals, 1):
            c = ws_manual.cell(row=ri, column=ci)
            c.value     = val
            c.fill      = fill
            c.border    = thin_border()
            c.alignment = center() if ci in [1, 3, 4, 5, 6, 8, 9] else left()
            c.font      = body_font(color=AMBER_DARK if ci == 9 and val == "No" else "000000",
                                    bold=(ci == 9 and val == "No"))
        ws_manual.row_dimensions[ri].height = 20

    wb.save(output_path)
    print(f"[+] Excel report written: {output_path}")


# -----------------------------------------------------------------------
# JSON Report
# -----------------------------------------------------------------------

def generate_json_report(records: list, output_path: str):
    data = {
        "report_generated": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "policy_version": "2024-Q1",
        "total_connections": len(records),
        "certified": sum(1 for r in records if r.recertification_status == "CERTIFIED"),
        "failed": sum(1 for r in records if r.recertification_status == "FAILED"),
        "manual_review": sum(1 for r in records if r.recertification_status == "MANUAL_REVIEW"),
        "connections": [asdict(r) for r in records],
    }
    with open(output_path, "w") as f:
        json.dump(data, f, indent=2)
    print(f"[+] JSON report written: {output_path}")


# -----------------------------------------------------------------------
# Console Summary
# -----------------------------------------------------------------------

def print_summary(records: list):
    print("\n" + "="*70)
    print("   THIRD-PARTY IN-TRANSIT ENCRYPTION RECERTIFICATION SUMMARY")
    print("="*70)

    external = [r for r in records if r.recertification_status != "MANUAL_REVIEW"]
    for region in ["AMER", "EMEA", "APAC"]:
        regional = [r for r in external if r.region == region]
        if not regional:
            continue
        certified = sum(1 for r in regional if r.recertification_status == "CERTIFIED")
        failed    = sum(1 for r in regional if r.recertification_status == "FAILED")
        print(f"\n  {region}: {len(regional)} connections | Certified: {certified} | Failed: {failed}")
        for r in regional:
            icon = "✓" if r.recertification_status == "CERTIFIED" else "✗"
            print(f"    [{icon}] {r.connection_id:10} | {r.vendor_name:35} | "
                  f"TLS: {r.tls_version or 'N/A':8} | {r.recertification_status}")
            for f in r.findings:
                if f != "PASS: All encryption checks passed":
                    print(f"           -> {f}")

    internal = [r for r in records if r.recertification_status == "MANUAL_REVIEW"]
    if internal:
        print(f"\n  {'='*66}")
        print(f"  INTERNAL ENDPOINTS — MANUAL REVIEW REQUIRED ({len(internal)} connections)")
        print(f"  {'='*66}")
        print(f"  Action: Escalate to Network Security for on-site TLS probe.")
        print()
        for r in internal:
            print(f"    [?] {r.connection_id:10} | {r.vendor_name:35} | "
                  f"{r.endpoint_host}:{r.endpoint_port}")
            for f in r.findings:
                print(f"           -> {f}")

    total     = len(records)
    certified = sum(1 for r in records if r.recertification_status == "CERTIFIED")
    failed    = sum(1 for r in records if r.recertification_status == "FAILED")
    manual    = sum(1 for r in records if r.recertification_status == "MANUAL_REVIEW")
    ext_total = total - manual
    print(f"\n  OVERALL EXTERNAL : {certified}/{ext_total} certified | "
          f"{failed}/{ext_total} failed "
          f"({100*certified//ext_total if ext_total else 0}% compliance rate)")
    print(f"  INTERNAL PENDING : {manual} connections awaiting Network Security review")
    print("="*70 + "\n")


# -----------------------------------------------------------------------
# Inventory
# -----------------------------------------------------------------------

SAMPLE_INVENTORY = [
    # --- Real banking vendors (CERTIFIED expected) ---
    ConnectionRecord("CONN-001", "Retail Banking API Team",   "Visa Inc.",              "visa.com",                     443,  "AMER", "PCI"),
    ConnectionRecord("CONN-002", "Corporate Banking",         "SWIFT Network",          "swift.com",                    443,  "AMER", "Confidential"),
    ConnectionRecord("CONN-003", "Digital Channels",          "Plaid Technologies",     "plaid.com",                    443,  "AMER", "PII"),
    ConnectionRecord("CONN-004", "FX Trading Platform",       "Bloomberg LP",           "bloomberg.com",                443,  "EMEA", "Confidential"),
    ConnectionRecord("CONN-005", "KYC/AML Compliance",        "LexisNexis Risk",        "lexisnexis.com",               443,  "EMEA", "PII"),
    ConnectionRecord("CONN-006", "Payments Gateway",          "Mastercard Intl.",       "mastercard.com",               443,  "APAC", "PCI"),
    ConnectionRecord("CONN-007", "Cloud Infrastructure",      "AWS Financial Svcs",     "amazonaws.com",                443,  "APAC", "Internal"),
    ConnectionRecord("CONN-008", "Credit Risk Engine",        "Moody's Analytics",      "moodys.com",                   443,  "AMER", "Confidential"),
    ConnectionRecord("CONN-009", "CRM Platform",              "Salesforce Inc.",        "salesforce.com",               443,  "AMER", "Confidential"),
    ConnectionRecord("CONN-010", "Retail Banking Portal",     "TD Bank",                "td.com",                       443,  "AMER", "PCI"),

    # --- BadSSL test cases (FAILED expected) ---
    ConnectionRecord("CONN-012", "Legacy Settlements",        "Finacle Core System",    "expired.badssl.com",           443,  "EMEA", "Confidential"),
    ConnectionRecord("CONN-013", "Trade Matching Engine",     "Temenos T24",            "wrong.host.badssl.com",        443,  "EMEA", "PCI"),
    ConnectionRecord("CONN-014", "Correspondent Banking API", "FIS Global Legacy",      "self-signed.badssl.com",       443,  "APAC", "Confidential"),
    ConnectionRecord("CONN-015", "Sanctions Screening Feed",  "Oracle FLEXCUBE",        "untrusted-root.badssl.com",    443,  "APAC", "PII"),
    ConnectionRecord("CONN-016", "Bond Trading Platform",     "Murex MX.3 Gateway",     "revoked.badssl.com",           443,  "EMEA", "Confidential"),
    ConnectionRecord("CONN-017", "Interbank Messaging",       "FiServ Legacy Adapter",  "incomplete-chain.badssl.com",  443,  "AMER", "Confidential"),
    ConnectionRecord("CONN-018", "Vendor Payment Bridge",     "Legacy FX Settlement",   "rc4.badssl.com",               443,  "AMER", "PCI"),
    ConnectionRecord("CONN-019", "Core Banking Interface",    "Finacle TLSv1.0 Node",   "tls-v1-0.badssl.com",         1010, "APAC", "Confidential"),
    ConnectionRecord("CONN-020", "Trade Finance Portal",      "Temenos TLSv1.1 Node",   "tls-v1-1.badssl.com",         1011, "EMEA", "PCI"),
    ConnectionRecord("CONN-021", "PKI Certificate Service",   "Entrust SHA1 Bridge",    "sha1-intermediate.badssl.com", 443,  "AMER", "Internal"),

    # --- Internal network endpoints (MANUAL_REVIEW expected) ---
    ConnectionRecord("CONN-022", "Core Banking Team",         "App Server 01 (LAB-PROD-AS01)",    "10.0.1.50",     8443, "AMER", "Confidential"),
    ConnectionRecord("CONN-023", "Payments Infrastructure",   "Load Balancer Primary (LB-PAY01)", "10.0.2.10",     443,  "AMER", "PCI"),
    ConnectionRecord("CONN-024", "Database Administration",   "Oracle DB Primary (DB-CORE-01)",   "10.0.5.12",     2484, "AMER", "PCI"),
    ConnectionRecord("CONN-025", "Network Security",          "Core Firewall MGMT (FW-CORE-01)",  "172.16.0.1",    8443, "AMER", "Internal"),
    ConnectionRecord("CONN-026", "FX Trading Infra",          "Trading Engine Node (EMEA-TX-01)", "172.16.10.55",  9443, "EMEA", "Confidential"),
    ConnectionRecord("CONN-027", "Integration Middleware",    "API Gateway Node (APAC-GW-01)",    "192.168.5.100", 9443, "APAC", "Internal"),
    ConnectionRecord("CONN-028", "KYC/AML Platform",          "Screening Server (AML-SRV-02)",    "192.168.10.22", 8080, "APAC", "PII"),
]


# -----------------------------------------------------------------------
# Main
# -----------------------------------------------------------------------

if __name__ == "__main__":
    print("[*] Starting in-transit encryption scan of third-party connections...\n")

    scanned = []
    for record in SAMPLE_INVENTORY:
        print(f"[*] Scanning {record.connection_id}: {record.vendor_name} "
              f"({record.endpoint_host}:{record.endpoint_port})")
        result = scan_endpoint(record)
        scanned.append(result)

    print_summary(scanned)

    ts = datetime.datetime.now(datetime.timezone.utc).strftime("%Y%m%d_%H%M%S")
    generate_excel_report(scanned, f"Encryption_Recertification_{ts}.xlsx")
    generate_json_report(scanned,  f"scan_{ts}.json")